"""Bounded uploads -> durable preview -> explicit review -> atomic commit.

No URL/path reads, fuzzy company creation, title upserts, or inferred completion.
This importer intentionally handles activities and their single imported session.
"""
import csv
from datetime import date, datetime, time, timezone
import hashlib
from io import BytesIO, StringIO
import json
from pathlib import PurePath
import re
from typing import Any
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from psycopg2.extras import Json, RealDictCursor
from pydantic import ValidationError

from core.db import nl_db
from core.exceptions import AppException, BadRequestError, NotFoundError
from schemas.friday_import import (
    ImportBatch, ImportCommit, ImportNormalized, ImportReceipt, ImportReview,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 30 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 40
MAX_CELL_LENGTH = 2000
HEADERS = (
    "sourceKey", "academicYear", "term", "title", "companyId", "companyName",
    "sessionDate", "startTime", "endTime", "location", "format", "capacity",
    "booked", "attended", "status",
)
ALIASES = {
    "ปีการศึกษา": "academicYear", "ภาคการศึกษา": "term", "เทอม": "term",
    "ชื่อกิจกรรม": "title", "หัวข้อ": "title", "รหัสบริษัท": "companyId",
    "บริษัท": "companyName", "ชื่อบริษัท": "companyName", "วันที่": "sessionDate",
    "เวลาเริ่ม": "startTime", "เวลาสิ้นสุด": "endTime", "สถานที่": "location",
    "รูปแบบ": "format", "รับสมัคร": "capacity", "จำนวนจอง": "booked",
    "ผู้เข้าร่วม": "attended", "สถานะ": "status", "รหัสต้นทาง": "sourceKey",
}
TERMS = {"1": "FIRST", "first": "FIRST", "2": "SECOND", "second": "SECOND", "summer": "SUMMER", "ฤดูร้อน": "SUMMER"}
FORMATS = {"talk", "workshop", "visit", "hackathon", "other", "unknown"}
STATUSES = {"draft", "contacting", "invited", "confirmed", "in_progress", "completed", "cancelled"}
STATUS_ALIASES = {"held": "completed", "scheduled": "draft", "จัดแล้ว": "completed", "ยกเลิก": "cancelled"}


def _conflict(message):
    raise AppException(message=message, status_code=409)


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str).encode()).hexdigest()


def _identifier(value):
    return isinstance(value, str) and re.fullmatch(r"[1-9][0-9]{0,18}", value) and int(value) <= 9223372036854775807


def _cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        value = value.date().isoformat() if value.time() == time() else value.isoformat()
    elif isinstance(value, (date, time)):
        value = value.isoformat(timespec="minutes") if isinstance(value, time) else value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        value = int(value)
    value = str(value).strip()
    if len(value) > MAX_CELL_LENGTH:
        raise BadRequestError(message=f"ช่องข้อมูลยาวเกิน {MAX_CELL_LENGTH} ตัวอักษร")
    if "\x00" in value:
        raise BadRequestError(message="ไฟล์มีอักขระที่ไม่รองรับ")
    return value or None


def _sheet_rows(values, sheet_name):
    iterator = iter(values)
    try:
        first = next(iterator)
    except StopIteration:
        return []
    if len(first) > MAX_COLUMNS:
        raise BadRequestError(message=f"ไฟล์ต้องมีไม่เกิน {MAX_COLUMNS} คอลัมน์")
    keys = []
    canonical = {h.casefold(): h for h in HEADERS}
    for value in first:
        label = _cell(value) or ""
        key = ALIASES.get(label, canonical.get(label.casefold(), label))
        if not key or key not in HEADERS or key in keys:
            raise BadRequestError(message=f"หัวตารางไม่รองรับหรือซ้ำ: {label[:80]}")
        keys.append(key)
    if not {"academicYear", "term", "title"}.issubset(keys):
        raise BadRequestError(message="หัวตารางต้องมี academicYear, term และ title")
    rows = []
    for number, values in enumerate(iterator, start=2):
        if number > MAX_ROWS + 1:
            raise BadRequestError(message=f"นำเข้าได้ครั้งละไม่เกิน {MAX_ROWS} แถว")
        if len(values) > len(keys) and any(_cell(v) is not None for v in values[len(keys):]):
            raise BadRequestError(message=f"แถว {number} มีข้อมูลเกินหัวตาราง")
        raw = {key: _cell(values[i]) if i < len(values) else None for i, key in enumerate(keys)}
        rows.append({"row_number": number, "sheet_name": sheet_name, "raw": raw,
                     "row_identity": _digest(raw)})
    return rows


def parse_upload(content: bytes, filename: str):
    if not content or len(content) > MAX_FILE_BYTES:
        raise BadRequestError(message=f"ไฟล์ต้องไม่ว่างและมีขนาดไม่เกิน {MAX_FILE_BYTES // 1024 // 1024} MB")
    extension = PurePath(filename).suffix.lower()
    try:
        if extension == ".csv":
            try:
                source = content.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise BadRequestError(message="CSV ต้องเข้ารหัส UTF-8") from exc
            return _sheet_rows(csv.reader(StringIO(source), strict=True), "CSV")
        if extension != ".xlsx":
            raise BadRequestError(message="รองรับเฉพาะไฟล์ .csv UTF-8 และ .xlsx")
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 300 or sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise BadRequestError(message="ไฟล์ Excel ขยายใหญ่เกินขนาดที่รองรับ")
            if any(item.flag_bits & 1 for item in archive.infolist()):
                raise BadRequestError(message="ไม่รองรับ Excel ที่เข้ารหัส")
            for item in archive.infolist():
                if item.filename.lower().endswith((".xml", ".rels")):
                    xml = archive.read(item).upper()
                    if b"\x00" in xml or b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise BadRequestError(message="ไม่รองรับ XML ที่ประกาศ DTD หรือ entity")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise BadRequestError(message="รุ่นนี้รองรับ Excel หนึ่งชีตต่อไฟล์ โปรดแยกชีตที่ต้องการนำเข้า")
            sheet = workbook.worksheets[0]
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                raise BadRequestError(message=f"นำเข้าได้ครั้งละไม่เกิน {MAX_ROWS} แถว")
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise BadRequestError(message=f"ไฟล์ต้องมีไม่เกิน {MAX_COLUMNS} คอลัมน์")
            return _sheet_rows(sheet.iter_rows(values_only=True), sheet.title)
        finally:
            workbook.close()
    except AppException:
        raise
    except Exception as exc:
        raise BadRequestError(message="อ่านไฟล์ไม่ได้ โปรดใช้แม่แบบ CSV หรือ Excel ที่ถูกต้อง") from exc


def _integer(raw, field, errors):
    value = raw.get(field)
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not re.fullmatch(r"[0-9]+", str(value)) or int(value) > 2147483647:
        errors.append(f"{field}: ต้องเป็นจำนวนเต็มตั้งแต่ 0 ถึง 2147483647")
        return None
    return int(value)


def normalize_row(raw, year_convention="AD"):
    """Return errors beside values, so invalid dates/counts never become facts."""
    errors = []
    raw = {key: _cell(value) for key, value in raw.items()}
    title = raw.get("title") or ""
    if not title:
        errors.append("title: ไม่มีชื่อกิจกรรม กรุณาเลือกข้ามแถวว่างอย่างชัดเจนหรือเพิ่มชื่อก่อนเลือกนำเข้า")
    if len(title) > 500:
        errors.append("title: ยาวเกิน 500 ตัวอักษร")
    year = _integer(raw, "academicYear", errors)
    if year is not None and year_convention == "BE":
        year -= 543
    if year is None or not 2000 <= year <= 2100:
        errors.append("academicYear: ระบุปี 2000-2100 ค.ศ. หรือเลือก BE สำหรับ 2543-2643 พ.ศ.")
    term = TERMS.get((raw.get("term") or "").casefold())
    if term is None:
        errors.append("term: รองรับ FIRST, SECOND, SUMMER หรือ 1, 2, summer")
    status = (raw.get("status") or "").casefold() or None
    status = STATUS_ALIASES.get(status, status)
    if status not in STATUSES:
        errors.append("status: ต้องระบุสถานะที่ทราบชัด ห้ามอนุมานจากวันที่")
    fmt = (raw.get("format") or "unknown").casefold()
    if fmt not in FORMATS:
        errors.append("format: เลือก talk, workshop, visit, hackathon, other หรือ unknown")
    when = raw.get("sessionDate")
    if when:
        try:
            if not re.fullmatch(r"20[0-9]{2}-[0-9]{2}-[0-9]{2}", when):
                raise ValueError()
            date.fromisoformat(when)
        except ValueError:
            errors.append("sessionDate: ต้องเป็นวันที่จริงรูปแบบ ค.ศ. YYYY-MM-DD")
    start, end = raw.get("startTime"), raw.get("endTime")
    for key, value in (("startTime", start), ("endTime", end)):
        if value and not re.fullmatch(r"(?:[01][0-9]|2[0-3]):[0-5][0-9]", value):
            errors.append(f"{key}: ใช้เวลา HH:MM เขต Asia/Bangkok")
    if start and end and end <= start:
        errors.append("endTime: ต้องหลัง startTime ในวันเดียวกัน")
    location = raw.get("location")
    if location and len(location) > 1000:
        errors.append("location: ยาวเกิน 1000 ตัวอักษร")
    if status in {"confirmed", "in_progress"} and not all((when, start, end, location)):
        errors.append("status: ยืนยัน/กำลังจัดต้องมีวัน เวลาเริ่ม เวลาสิ้นสุด และสถานที่")
    counts = {key: _integer(raw, key, errors) for key in ("capacity", "booked", "attended")}
    source_key = raw.get("sourceKey")
    if source_key and len(source_key) > 200:
        errors.append("sourceKey: ยาวเกิน 200 ตัวอักษร")
    normalized = ImportNormalized(source_key=source_key, academic_year=year, term=term,
        title=title, company_name=raw.get("companyName") or "", session_date=when,
        start_time=start, end_time=end, location=location, format=fmt, status=status, **counts)
    return normalized, errors


def _companies(cursor, raw, explicit_ids=None):
    if explicit_ids is not None:
        ids = list(dict.fromkeys(explicit_ids))
        if not ids or any(not _identifier(value) for value in ids):
            return [], "needs_mapping", ["companyIds: เลือกบริษัทที่มีอยู่ในทะเบียนอย่างน้อยหนึ่งบริษัท"]
        cursor.execute("SELECT id FROM companies WHERE id = ANY(%s::bigint[]) ORDER BY id", (ids,))
        found = [str(row["id"]) for row in cursor.fetchall()]
        if set(found) != set(ids):
            return [], "needs_mapping", ["companyIds: ไม่พบบางบริษัทในทะเบียน กรุณาจับคู่ใหม่"]
        return sorted(ids, key=int), "matched", []
    raw_id = raw.get("companyId")
    if raw_id:
        return _companies(cursor, raw, [raw_id])
    name = raw.get("companyName")
    if name:
        cursor.execute("""SELECT id FROM companies WHERE lower(btrim(company_th))=lower(%s)
            OR lower(btrim(company_en))=lower(%s) OR EXISTS (
            SELECT 1 FROM unnest(aliases) AS alias WHERE lower(btrim(alias))=lower(%s)) ORDER BY id""", (name, name, name))
        found = cursor.fetchall()
        if len(found) == 1:
            return [str(found[0]["id"])], "matched", []
    return [], "needs_mapping", ["company: ไม่พบชื่อที่ตรงเพียงบริษัทเดียว กรุณาเลือกบริษัทจากทะเบียน"]


def _candidates(cursor, normalized):
    cursor.execute("""SELECT DISTINCT a.id, a.title, a.revision FROM friday_activities a
        JOIN academic_periods p ON p.id=a.period_id
        WHERE (p.academic_year=%s AND p.term=%s AND lower(btrim(a.title))=lower(%s))
        OR EXISTS (SELECT 1 FROM friday_import_source_links s WHERE s.activity_id=a.id
          AND %s IS NOT NULL AND s.source_key=%s)
        ORDER BY a.id LIMIT 20""", (normalized.academic_year, normalized.term, normalized.title,
                                    normalized.source_key, normalized.source_key))
    return [{"id": str(row["id"]), "title": row["title"], "revision": row["revision"]} for row in cursor.fetchall()]


def _audit(cursor, batch_id, actor, operation, details, row_id=None):
    cursor.execute("""INSERT INTO friday_import_audit(batch_id,row_id,actor_user_id,operation,details)
        VALUES(%s,%s,%s,%s,%s)""", (batch_id, row_id, actor, operation, Json(details)))


def _batch(cursor, batch_id, *, lock=False):
    cursor.execute("SELECT * FROM friday_import_batches WHERE id=%s" + (" FOR UPDATE" if lock else ""), (batch_id,))
    batch = cursor.fetchone()
    if batch is None:
        raise NotFoundError(message="ไม่พบชุดนำเข้า")
    return batch


def _result(cursor, batch_id, receipt=None, replay=False):
    batch = _batch(cursor, batch_id)
    cursor.execute("SELECT * FROM friday_import_rows WHERE batch_id=%s ORDER BY id", (batch_id,))
    rows = []
    for row in cursor.fetchall():
        rows.append({"id": str(row["id"]), "row_number": row["row_number"], "sheet_name": row["sheet_name"],
            "raw": row["raw"], "normalized": row["normalized"], "errors": row["errors"],
            "mapping_status": row["mapping_status"], "candidate_activities": row["candidates"],
            "decision": row["decision"], "target_activity_id": str(row["target_activity_id"]) if row["target_activity_id"] else None,
            "target_activity_revision": row["target_activity_revision"],
            "committed_activity_id": str(row["committed_activity_id"]) if row["committed_activity_id"] else None})
    if receipt is None:
        cursor.execute("SELECT receipt FROM friday_import_receipts WHERE batch_id=%s ORDER BY committed_at DESC LIMIT 1", (batch_id,))
        previous = cursor.fetchone()
        receipt = previous["receipt"] if previous else None
    return ImportBatch(id=str(batch["id"]), filename=batch["filename"], revision=batch["revision"],
        status=batch["status"], year_convention=batch["year_convention"], source_hash=batch["source_hash"],
        rows=rows, receipt=receipt, idempotent_replay=replay)


def preview_import(content, filename, year_convention, actor_user_id):
    if year_convention not in {"AD", "BE"}:
        raise BadRequestError(message="yearConvention ต้องเป็น AD หรือ BE")
    source_rows = parse_upload(content, filename)
    if not source_rows:
        raise BadRequestError(message="ไฟล์ไม่มีแถวข้อมูล")
    batch_id = str(uuid4())
    safe_name = str(filename).replace("\\", "/").split("/")[-1][:255]
    with nl_db.get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cursor:
            cursor.execute("""INSERT INTO friday_import_batches(id,actor_user_id,filename,source_hash,year_convention)
                VALUES(%s,%s,%s,%s,%s)""", (batch_id, actor_user_id, safe_name, hashlib.sha256(content).hexdigest(), year_convention))
            for row in source_rows:
                normalized, errors = normalize_row(row["raw"], year_convention)
                company_ids, mapping, mapping_errors = _companies(cursor, row["raw"])
                normalized.company_ids = company_ids
                errors += mapping_errors
                cursor.execute("""INSERT INTO friday_import_rows(batch_id,row_number,sheet_name,row_identity,raw,
                    normalized,errors,mapping_status,candidates,decision) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (batch_id, row["row_number"], row["sheet_name"], row["row_identity"], Json(row["raw"]),
                     Json(normalized.model_dump(mode="json", by_alias=True)), Json(errors), mapping,
                     Json(_candidates(cursor, normalized)), "pending"))
            _audit(cursor, batch_id, actor_user_id, "preview", {"rowCount": len(source_rows), "yearConvention": year_convention})
            result = _result(cursor, batch_id)
        conn.commit()
    return result


def get_import(batch_id):
    with nl_db.get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cursor:
            return _result(cursor, batch_id)


def review_import_row(batch_id, row_id, payload: ImportReview, actor_user_id):
    if not _identifier(row_id):
        raise BadRequestError(message="รหัสแถวนำเข้าไม่ถูกต้อง")
    with nl_db.get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cursor:
            batch = _batch(cursor, batch_id, lock=True)
            if batch["revision"] != payload.expected_batch_revision:
                _conflict("ชุดนำเข้าถูกแก้ไขแล้ว กรุณาโหลดข้อมูลล่าสุดก่อนตรวจแถวอีกครั้ง")
            cursor.execute("SELECT * FROM friday_import_rows WHERE batch_id=%s AND id=%s", (batch_id, row_id))
            row = cursor.fetchone()
            if row is None:
                raise NotFoundError(message="ไม่พบแถวนำเข้า")
            if row["committed_activity_id"]:
                _conflict("แถวนี้นำเข้าแล้ว กรุณาแก้ข้อมูลจากหน้ากิจกรรม")
            current = ImportNormalized.model_validate(row["normalized"])
            # Preserve invalid raw values until explicitly corrected: a failed
            # count parse must not turn into an apparently valid unknown count.
            values = dict(row["raw"])
            if batch["year_convention"] == "BE" and values.get("academicYear") and re.fullmatch(r"[0-9]+", values["academicYear"]):
                values["academicYear"] = str(int(values["academicYear"]) - 543)
            corrections = dict(row["corrections"])
            if payload.normalized:
                corrections.update(payload.normalized.model_dump(mode="json", by_alias=True, exclude_unset=True))
            values.update(corrections)
            normalized, errors = normalize_row(values, "AD")
            normalized.company_name = current.company_name
            ids, mapping, mapping_errors = _companies(cursor, row["raw"], payload.company_ids)
            normalized.company_ids = ids
            errors += mapping_errors
            if payload.decision == "update":
                if not _identifier(payload.target_activity_id) or payload.target_activity_revision is None:
                    raise BadRequestError(message="การอัปเดตต้องเลือก targetActivityId และ targetActivityRevision")
                cursor.execute("SELECT id,status,revision FROM friday_activities WHERE id=%s", (payload.target_activity_id,))
                target = cursor.fetchone()
                if target is None:
                    raise BadRequestError(message="ไม่พบกิจกรรมที่ต้องการอัปเดต")
                if target["revision"] != payload.target_activity_revision:
                    errors.append("targetActivityRevision: กิจกรรมถูกแก้ไขแล้ว กรุณาเลือกเป้าหมายรุ่นล่าสุด")
                cursor.execute("SELECT count(*) AS n FROM friday_sessions WHERE activity_id=%s", (payload.target_activity_id,))
                if cursor.fetchone()["n"] > 1:
                    errors.append("targetActivityId: รุ่นนี้อัปเดตผ่านไฟล์ได้เฉพาะกิจกรรมที่มีไม่เกินหนึ่งรอบ กรุณาแก้กิจกรรมหลายรอบจากหน้ารายละเอียด")
                errors += _target_errors(target["status"], normalized)
            elif payload.target_activity_id is not None or payload.target_activity_revision is not None:
                raise BadRequestError(message="ระบุกิจกรรมเป้าหมายได้เฉพาะการอัปเดต")
            cursor.execute("""UPDATE friday_import_rows SET normalized=%s,corrections=%s,errors=%s,mapping_status=%s,
                candidates=%s,decision=%s,target_activity_id=%s,target_activity_revision=%s,reviewed_by=%s,reviewed_at=now()
                WHERE id=%s""", (Json(normalized.model_dump(mode="json", by_alias=True)), Json(corrections), Json(errors), mapping,
                Json(_candidates(cursor, normalized)), payload.decision, payload.target_activity_id,
                payload.target_activity_revision, actor_user_id, row_id))
            cursor.execute("UPDATE friday_import_batches SET revision=revision+1,status='preview',updated_at=now() WHERE id=%s", (batch_id,))
            _audit(cursor, batch_id, actor_user_id, "review", {"decision": payload.decision,
                "targetActivityId": payload.target_activity_id, "normalized": normalized.model_dump(mode="json", by_alias=True)}, row_id)
            result = _result(cursor, batch_id)
        conn.commit()
    return result


def _write_row(conn, row, normalized, actor_user_id, batch_id):
    from schemas.friday_activity import FridayActivityCreate, FridayActivityUpdate, FridaySessionInput, FridayTransition
    from services.friday_activity import (
        _activity, create_activity_in_transaction, update_activity_in_transaction,
        transition_activity_in_transaction,
    )
    reason = f"Reviewed Friday import {batch_id}, row {row['row_number']}"
    session = FridaySessionInput(session_date=normalized.session_date, start_time=normalized.start_time,
        end_time=normalized.end_time, location=normalized.location,
        status="held" if normalized.status == "completed" else "cancelled" if normalized.status == "cancelled" else "planned",
        capacity=normalized.capacity, booked=normalized.booked, attended=normalized.attended)
    fields = {"academic_year": normalized.academic_year, "term": normalized.term, "title": normalized.title,
              "format": normalized.format, "company_ids": normalized.company_ids, "sessions": [session]}
    if row["decision"] == "create":
        fields["owner_user_id"] = str(actor_user_id)
        activity = create_activity_in_transaction(conn,
            FridayActivityCreate(**fields, idempotency_key=f"{batch_id}:{row['id']}"), str(actor_user_id), operation="import.create")
    else:
        before = _activity(conn, str(row["target_activity_id"]), lock=True)
        if before.revision != row["target_activity_revision"]:
            _conflict(f"กิจกรรมเป้าหมายของแถว {row['row_number']} ถูกแก้ไขแล้ว กรุณาตรวจแถวใหม่")
        if len(before.sessions) > 1:
            raise BadRequestError(message="ไม่สามารถนำเข้าทับกิจกรรมหลายรอบได้ กรุณาแก้จากหน้ารายละเอียด")
        if before.sessions:
            session.id = before.sessions[0].id
        fields.update(owner_user_id=before.owner_user_id, description=before.description,
                      domain=before.domain, people=before.people)
        reopening_completed = before.status == "completed" and normalized.status == "confirmed"
        if reopening_completed:
            # Keep the held evidence while patching a closed record. Reopen
            # after its new schedule is complete, then mark the new plan. Every
            # phase is audited and shares the caller's single transaction.
            session.status = "held"
        activity = update_activity_in_transaction(conn, before.id,
            FridayActivityUpdate(**fields, expected_revision=before.revision, reason=reason), str(actor_user_id), reason=reason)
        if reopening_completed:
            activity = transition_activity_in_transaction(conn, activity.id,
                FridayTransition(expected_revision=activity.revision, status="confirmed", reason=reason),
                str(actor_user_id), operation="import.transition")
            session.status = "planned"
            return update_activity_in_transaction(conn, activity.id,
                FridayActivityUpdate(**fields, expected_revision=activity.revision, reason=reason), str(actor_user_id), reason=reason)
    desired = normalized.status
    # in_progress is explicit in the reviewed source; record each legal step.
    if ((desired == "in_progress" and activity.status in {"draft", "contacting", "invited"})
        or (desired == "completed" and activity.status in {"contacting", "invited"})):
        activity = transition_activity_in_transaction(conn, activity.id,
            FridayTransition(expected_revision=activity.revision, status="confirmed", reason=reason), str(actor_user_id), operation="import.transition")
    if desired == "completed" and activity.status == "confirmed":
        activity = transition_activity_in_transaction(conn, activity.id,
            FridayTransition(expected_revision=activity.revision, status="in_progress", reason=reason), str(actor_user_id), operation="import.transition")
    if desired != activity.status:
        activity = transition_activity_in_transaction(conn, activity.id,
            FridayTransition(expected_revision=activity.revision, status=desired, reason=reason), str(actor_user_id),
            allow_historical=True, operation="import.transition")
    return activity


def _target_errors(current_status, normalized):
    """Show known lifecycle incompatibilities during review, before commit."""
    from services.friday_activity import _FORWARD
    desired = normalized.status
    if desired not in STATUSES:
        return []  # Its source validation already explains the unknown status.
    reopening = (current_status, desired) in {("completed", "confirmed"), ("cancelled", "draft")}
    if current_status in {"completed", "cancelled"} and desired != current_status and not reopening:
        return ["status: การเปิดงานที่ปิดแล้วหรือเปลี่ยนสถานะประวัติต้องทำจากหน้ากิจกรรมก่อนนำเข้า"]
    permitted = (desired == current_status or desired in _FORWARD[current_status]
        or (current_status == "draft" and desired in {"completed", "in_progress"})
        or (current_status in {"contacting", "invited"} and desired in {"completed", "in_progress"})
        or (current_status == "confirmed" and desired == "completed"))
    errors = []
    if not permitted:
        errors.append("status: ไม่สามารถเปลี่ยนจากสถานะกิจกรรมเป้าหมายเป็นสถานะในไฟล์ กรุณาแก้จากหน้ากิจกรรมก่อน")
    if (current_status in {"confirmed", "in_progress"}
        or (desired == "completed" and current_status in {"contacting", "invited"})):
        if not all((normalized.session_date, normalized.start_time, normalized.end_time, normalized.location)):
            errors.append("status: การอัปเดตงานที่ยืนยันแล้วต้องมีวัน เวลาเริ่ม เวลาสิ้นสุด และสถานที่ครบ")
    return errors


def commit_import(batch_id, payload: ImportCommit, actor_user_id):
    if not payload.idempotency_key.strip():
        raise BadRequestError(message="idempotencyKey ต้องไม่ว่าง")
    with nl_db.get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cursor:
            cursor.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s,0))",
                           (f"friday-import-key:{actor_user_id}:{payload.idempotency_key}",))
            cursor.execute("""SELECT batch_id,review_revision,receipt FROM friday_import_receipts
                WHERE actor_user_id=%s AND idempotency_key=%s""", (actor_user_id, payload.idempotency_key))
            previous = cursor.fetchone()
            if previous:
                if str(previous["batch_id"]) != batch_id or previous["review_revision"] != payload.expected_batch_revision:
                    _conflict("idempotencyKey นี้ถูกใช้กับชุดหรือรุ่นนำเข้าอื่นแล้ว")
                result = _result(cursor, batch_id, receipt=previous["receipt"], replay=True)
                conn.rollback()
                return result
            batch = _batch(cursor, batch_id, lock=True)
            if batch["revision"] != payload.expected_batch_revision:
                _conflict("ชุดนำเข้าถูกแก้ไขแล้ว กรุณาโหลดข้อมูลล่าสุดก่อนนำเข้า")
            # Serialize different batches of the same exact file before probing
            # their source links, including concurrently submitted first copies.
            cursor.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s,0))", (f"friday-import-file:{batch['source_hash']}",))
            cursor.execute("SELECT * FROM friday_import_rows WHERE batch_id=%s ORDER BY id", (batch_id,))
            rows = cursor.fetchall()
            selected = [r for r in rows if r["decision"] in {"create", "update"} and not r["committed_activity_id"]]
            if any(r["errors"] or r["mapping_status"] != "matched" or r["reviewed_by"] is None for r in selected):
                raise BadRequestError(message="แถวที่เลือกนำเข้ายังมีข้อผิดพลาดหรือยังไม่ตรวจ กรุณาแก้หรือเลือกข้ามก่อน")
            if not selected and not any(r["decision"] == "skip" for r in rows):
                raise BadRequestError(message="ยังไม่มีแถวที่ตรวจและเลือกนำเข้า")
            selected_hash = _digest([{k: r[k] for k in ("id", "normalized", "decision", "target_activity_id", "target_activity_revision")} for r in selected])
            inserted = updated = unchanged = 0
            activity_ids = []
            for row in selected:
                normalized = ImportNormalized.model_validate(row["normalized"])
                ids, mapping, errors = _companies(cursor, {}, normalized.company_ids)
                if errors or mapping != "matched":
                    raise BadRequestError(message=f"แถว {row['row_number']}: บริษัทเปลี่ยนไป กรุณาจับคู่ใหม่")
                normalized.company_ids = ids
                mapping_hash = _digest(normalized.model_dump(mode="json", by_alias=True))
                cursor.execute("""SELECT activity_id FROM friday_import_source_links
                    WHERE source_hash=%s AND row_identity=%s AND mapping_hash=%s""",
                    (batch["source_hash"], row["row_identity"], mapping_hash))
                linked = cursor.fetchone()
                if linked:
                    activity_id = str(linked["activity_id"])
                    if row["decision"] == "update" and activity_id != str(row["target_activity_id"]):
                        _conflict(f"แถว {row['row_number']} มี provenance ผูกกิจกรรมอื่นแล้ว กรุณาตรวจเป้าหมาย")
                    unchanged += 1
                else:
                    try:
                        activity = _write_row(conn, row, normalized, actor_user_id, batch_id)
                    except ValidationError as exc:
                        raise BadRequestError(message=f"แถว {row['row_number']}: ข้อมูลกิจกรรมไม่ผ่านการตรวจ") from exc
                    activity_id = activity.id
                    inserted += row["decision"] == "create"
                    updated += row["decision"] == "update"
                    cursor.execute("""INSERT INTO friday_import_source_links(source_hash,row_identity,mapping_hash,source_key,
                        activity_id,import_row_id,actor_user_id) VALUES(%s,%s,%s,%s,%s,%s,%s)""",
                        (batch["source_hash"], row["row_identity"], mapping_hash, normalized.source_key,
                         activity_id, row["id"], actor_user_id))
                cursor.execute("UPDATE friday_import_rows SET committed_activity_id=%s WHERE id=%s", (activity_id, row["id"]))
                activity_ids.append(activity_id)
            receipt = ImportReceipt(id=str(uuid4()), batch_id=batch_id, inserted=inserted, updated=updated,
                unchanged=unchanged, skipped=sum(r["decision"] == "skip" for r in rows),
                pending=sum(r["decision"] == "pending" and not r["errors"] for r in rows),
                rejected=sum(r["decision"] == "pending" and bool(r["errors"]) for r in rows),
                activity_ids=list(dict.fromkeys(activity_ids)), committed_at=datetime.now(timezone.utc))
            encoded = receipt.model_dump(mode="json", by_alias=True)
            cursor.execute("""INSERT INTO friday_import_receipts(id,batch_id,actor_user_id,idempotency_key,
                review_revision,selected_hash,receipt) VALUES(%s,%s,%s,%s,%s,%s,%s)""",
                (receipt.id, batch_id, actor_user_id, payload.idempotency_key, payload.expected_batch_revision, selected_hash, Json(encoded)))
            remaining = any(r["decision"] == "pending" and not r["committed_activity_id"] for r in rows)
            cursor.execute("UPDATE friday_import_batches SET status=%s,revision=revision+1,updated_at=now() WHERE id=%s",
                           ("preview" if remaining else "committed", batch_id))
            _audit(cursor, batch_id, actor_user_id, "commit", {"receiptId": receipt.id, "selectedHash": selected_hash,
                "reviewRevision": payload.expected_batch_revision, "inserted": inserted, "updated": updated, "unchanged": unchanged})
            result = _result(cursor, batch_id, receipt=encoded)
        conn.commit()
    return result
