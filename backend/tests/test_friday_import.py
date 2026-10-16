"""Real disposable PostgreSQL checks for provenance, atomic review and retries."""
from concurrent.futures import ThreadPoolExecutor
import csv
from io import BytesIO, StringIO
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
import psycopg2
import pytest

from conftest import TEST_URL
from core.exception_handlers import app_exception_handler
from core.exceptions import AppException, BadRequestError
from schemas.friday_import import ImportCommit
from services import friday_import as service

PREFIX = "/api/v1/friday-activities/imports"


def execute(sql, params=(), fetch=False):
    with psycopg2.connect(TEST_URL) as conn:
        with conn.cursor() as cursor:
            cursor.execute(sql, params)
            return cursor.fetchall() if fetch else None


@pytest.fixture
def import_client():
    from api.v1.auth import router as auth_router
    from api.v1.friday_import import router as import_router
    from api.v1.friday_activity import router as activity_router
    execute("UPDATE users SET role='staff'; INSERT INTO companies(company_th) VALUES ('Synthetic Labs');")
    app = FastAPI()
    app.add_exception_handler(AppException, app_exception_handler)
    app.include_router(auth_router, prefix="/api/v1")
    app.include_router(import_router, prefix="/api/v1")
    app.include_router(activity_router, prefix="/api/v1")
    with TestClient(app) as client:
        response = client.post("/api/v1/auth/login", json={"username": "alice", "password": "correct-password"})
        assert response.status_code == 200, response.text
        headers = {"Authorization": "Bearer " + response.json()["accessToken"]}
        yield client, headers


def csv_bytes(rows):
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=service.HEADERS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return output.getvalue().encode()


def source_row(**changes):
    return {"sourceKey": "synthetic-1", "academicYear": "2026", "term": "1", "title": "Synthetic activity",
            "companyId": "1", "sessionDate": "2026-10-16", "startTime": "13:00", "endTime": "16:00",
            "location": "Synthetic room", "format": "workshop", "capacity": "20", "booked": "0",
            "attended": "", "status": "draft", **changes}


def preview(client, headers, rows=None, *, content=None, name="synthetic.csv", convention="AD"):
    response = client.post(PREFIX + "/preview", headers=headers,
        files={"file": (name, content if content is not None else csv_bytes(rows or [source_row()]))},
        data={"yearConvention": convention})
    assert response.status_code == 201, response.text
    return response.json()


def review(client, headers, batch, index=0, **changes):
    row = batch["rows"][index]
    response = client.patch(f"{PREFIX}/{batch['id']}/rows/{row['id']}", headers=headers,
        json={"expectedBatchRevision": batch["revision"], "decision": "create", "companyIds": ["1"], **changes})
    assert response.status_code == 200, response.text
    return response.json()


def commit(client, headers, batch, key=None):
    response = client.post(f"{PREFIX}/{batch['id']}/commit", headers=headers,
        json={"expectedBatchRevision": batch["revision"], "idempotencyKey": key or str(uuid4())})
    assert response.status_code == 200, response.text
    return response.json()


def count(table):
    # Table names are constants within these synthetic tests.
    return execute(f"SELECT count(*) FROM {table}", fetch=True)[0][0]


def test_preview_is_durable_provenance_without_business_writes_and_blank_rows_visible(import_client):
    client, headers = import_client
    batch = preview(client, headers, [source_row(), source_row(title="", sourceKey="blank")])
    assert count("friday_activities") == 0
    assert count("friday_import_rows") == 2
    assert batch["rows"][0]["decision"] == "pending"
    assert batch["rows"][1]["decision"] == "pending"
    assert "ไม่มีชื่อกิจกรรม" in " ".join(batch["rows"][1]["errors"])
    assert client.get(f"{PREFIX}/{batch['id']}", headers=headers).json() == batch
    batch = review(client, headers, batch, 1, decision="skip", companyIds=[])
    assert batch["rows"][1]["decision"] == "skip"


def test_unknown_and_ambiguous_company_require_explicit_mapping(import_client):
    client, headers = import_client
    execute("INSERT INTO companies(company_th,company_en) VALUES ('Synthetic Other','Synthetic Labs')")
    batch = preview(client, headers, [source_row(companyId="", companyName="Synthetic Labs")])
    assert batch["rows"][0]["mappingStatus"] == "needs_mapping"
    assert batch["rows"][0]["normalized"]["companyIds"] == []
    batch = review(client, headers, batch)
    assert batch["rows"][0]["mappingStatus"] == "matched"
    assert commit(client, headers, batch)["receipt"]["inserted"] == 1
    assert count("companies") == 2


def test_invalid_values_persist_through_mapping_until_explicit_correction(import_client):
    client, headers = import_client
    batch = preview(client, headers, [source_row(companyId="", companyName="Unknown", attended="2.5",
        sessionDate="2026-02-30", format="magic", status="mystery")])
    batch = review(client, headers, batch)
    errors = " ".join(batch["rows"][0]["errors"])
    for field in ("attended", "sessionDate", "format", "status"):
        assert field in errors
    rejected = client.post(f"{PREFIX}/{batch['id']}/commit", headers=headers,
        json={"expectedBatchRevision": batch["revision"], "idempotencyKey": "bad-values"})
    assert rejected.status_code == 400
    assert count("friday_activities") == 0
    batch = review(client, headers, batch, normalized={"attended": None, "sessionDate": None, "format": "other", "status": "completed"})
    assert batch["rows"][0]["errors"] == []
    result = commit(client, headers, batch)
    activity = client.get("/api/v1/friday-activities/" + result["receipt"]["activityIds"][0], headers=headers).json()
    assert activity["status"] == "completed"
    assert activity["sessions"][0]["attended"] is None
    assert activity["sessions"][0]["sessionDate"] is None
    assert activity["sessions"][0]["booked"] == 0


def test_explicit_be_year_and_historical_status_no_date_based_completion(import_client):
    client, headers = import_client
    batch = preview(client, headers, [source_row(academicYear="2569", sessionDate="2020-01-01", attended="0", status="held")], convention="BE")
    assert batch["rows"][0]["normalized"]["academicYear"] == 2026
    batch = review(client, headers, batch)
    result = commit(client, headers, batch)
    assert execute("SELECT status FROM friday_activities", fetch=True) == [("completed",)]
    assert execute("SELECT status,attended FROM friday_sessions", fetch=True) == [("held", 0)]
    other = preview(client, headers, [source_row(sourceKey="past-draft", title="Past draft", sessionDate="2020-01-01")])
    commit(client, headers, review(client, headers, other))
    assert execute("SELECT status FROM friday_activities WHERE title='Past draft'", fetch=True) == [("draft",)]


def test_exact_file_mapping_and_retry_do_not_duplicate_business_records(import_client):
    client, headers = import_client
    staged = review(client, headers, preview(client, headers))
    result = commit(client, headers, staged, key="same-request")
    retry = commit(client, headers, staged, key="same-request")
    assert retry["idempotentReplay"] is True
    assert retry["receipt"] == result["receipt"]
    copied = review(client, headers, preview(client, headers))
    second = commit(client, headers, copied)
    assert second["receipt"]["inserted"] == 0
    assert second["receipt"]["unchanged"] == 1
    assert second["receipt"]["activityIds"] == result["receipt"]["activityIds"]
    assert count("friday_activities") == count("friday_sessions") == count("friday_import_source_links") == 1


def test_reordered_new_snapshot_stays_pending_and_explicit_target_updates_once(import_client):
    client, headers = import_client
    first_rows = [source_row(), source_row(sourceKey="synthetic-2", title="Second synthetic")]
    batch = preview(client, headers, first_rows)
    batch = review(client, headers, batch, 0)
    batch = review(client, headers, batch, 1)
    commit(client, headers, batch)
    snapshot = preview(client, headers, list(reversed(first_rows)))
    assert all(row["decision"] == "pending" for row in snapshot["rows"])
    assert all(len(row["candidateActivities"]) == 1 for row in snapshot["rows"])
    assert count("friday_activities") == 2
    for index in range(2):
        candidate = snapshot["rows"][index]["candidateActivities"][0]
        snapshot = review(client, headers, snapshot, index, decision="update",
            targetActivityId=candidate["id"], targetActivityRevision=candidate["revision"])
    result = commit(client, headers, snapshot)
    assert result["receipt"]["updated"] == 2
    assert count("friday_activities") == 2


def test_whole_selected_set_rolls_back_business_provenance_audit_receipt_then_retry(import_client, monkeypatch):
    client, headers = import_client
    batch = preview(client, headers, [source_row(), source_row(sourceKey="two", title="Second")])
    batch = review(client, headers, batch, 0)
    batch = review(client, headers, batch, 1)
    original = service._write_row
    calls = []
    def fail_second(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise AppException(message="Synthetic rollback probe", status_code=500)
        return original(*args, **kwargs)
    monkeypatch.setattr(service, "_write_row", fail_second)
    failed = client.post(f"{PREFIX}/{batch['id']}/commit", headers=headers,
        json={"expectedBatchRevision": batch["revision"], "idempotencyKey": "atomic-retry"})
    assert failed.status_code == 500
    for table in ("friday_activities", "friday_sessions", "friday_import_source_links", "friday_import_receipts", "friday_audit_events", "friday_mutation_receipts"):
        assert count(table) == 0
    assert execute("SELECT count(*) FROM friday_import_audit WHERE operation='commit'", fetch=True) == [(0,)]
    assert client.get(f"{PREFIX}/{batch['id']}", headers=headers).json()["revision"] == batch["revision"]
    monkeypatch.setattr(service, "_write_row", original)
    assert commit(client, headers, batch, key="atomic-retry")["receipt"]["inserted"] == 2


def test_concurrent_exact_file_batches_serialize_to_one_activity(import_client):
    client, headers = import_client
    first = review(client, headers, preview(client, headers))
    second = review(client, headers, preview(client, headers))
    def submit(batch):
        return service.commit_import(batch["id"], ImportCommit(expectedBatchRevision=batch["revision"], idempotencyKey=str(uuid4())), "1")
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(submit, (first, second)))
    assert sorted(result.receipt.inserted for result in results) == [0, 1]
    assert count("friday_activities") == 1


def test_review_and_target_revision_conflicts_preserve_existing_data(import_client):
    client, headers = import_client
    batch = review(client, headers, preview(client, headers))
    stale = client.patch(f"{PREFIX}/{batch['id']}/rows/{batch['rows'][0]['id']}", headers=headers,
        json={"expectedBatchRevision": batch["revision"] - 1, "decision": "skip"})
    assert stale.status_code == 409
    first = commit(client, headers, batch)
    activity_id = first["receipt"]["activityIds"][0]
    snapshot = preview(client, headers, [source_row(attended="3")])
    snapshot = review(client, headers, snapshot, decision="update", targetActivityId=activity_id, targetActivityRevision=1)
    execute("UPDATE friday_activities SET revision=revision+1 WHERE id=%s", (activity_id,))
    response = client.post(f"{PREFIX}/{snapshot['id']}/commit", headers=headers,
        json={"expectedBatchRevision": snapshot["revision"], "idempotencyKey": "stale-target"})
    assert response.status_code == 409
    assert execute("SELECT attended FROM friday_sessions", fetch=True) == [(None,)]
    assert count("friday_import_receipts") == 1


def test_duplicate_idempotency_key_cannot_retarget_batch(import_client):
    client, headers = import_client
    first = review(client, headers, preview(client, headers))
    commit(client, headers, first, key="bound-key")
    second = review(client, headers, preview(client, headers))
    response = client.post(f"{PREFIX}/{second['id']}/commit", headers=headers,
        json={"expectedBatchRevision": second["revision"], "idempotencyKey": "bound-key"})
    assert response.status_code == 409
    assert count("friday_activities") == 1


def test_partial_commit_receipt_keeps_unreviewed_rows_available_for_next_set(import_client):
    client, headers = import_client
    batch = preview(client, headers, [source_row(), source_row(sourceKey="pending", title="Pending corrected", attended="bad")])
    batch = review(client, headers, batch, 0)
    partial = commit(client, headers, batch, key="partial-first")
    assert partial["status"] == "preview"
    assert partial["receipt"]["inserted"] == 1
    assert partial["receipt"]["rejected"] == 1
    assert partial["rows"][0]["committedActivityId"] is not None
    assert partial["rows"][1]["committedActivityId"] is None
    corrected = review(client, headers, partial, 1, normalized={"attended": 0})
    result = commit(client, headers, corrected, key="partial-second")
    assert result["status"] == "committed"
    assert result["receipt"]["inserted"] == 1
    assert result["receipt"]["pending"] == result["receipt"]["rejected"] == 0
    assert count("friday_activities") == 2
    assert count("friday_import_receipts") == 2


def create_target(client, headers, *, sessions=1, status="draft"):
    body = {"academicYear": 2026, "term": "FIRST", "title": "Synthetic activity", "ownerUserId": "2",
        "description": "Preserve description", "domain": "Preserve domain", "format": "talk", "companyIds": ["1"],
        "people": [{"namespace": "guest", "name": "Synthetic coordinator", "role": "coordinator"}],
        "sessions": [{"sessionDate": "2026-10-16", "startTime": "13:00", "endTime": "16:00",
                      "location": "Original room", "status": "planned"} for _ in range(sessions)], "idempotencyKey": str(uuid4())}
    result = client.post("/api/v1/friday-activities", headers=headers, json=body)
    assert result.status_code == 201, result.text
    activity = result.json()
    if status != "draft":
        response = client.post(f"/api/v1/friday-activities/{activity['id']}/transitions", headers=headers,
            json={"expectedRevision": activity["revision"], "status": status})
        assert response.status_code == 200, response.text
        activity = response.json()
    return activity


def test_update_preserves_non_imported_metadata_and_single_session_identity(import_client):
    client, headers = import_client
    before = create_target(client, headers)
    batch = preview(client, headers, [source_row(attended="3", location="Updated room")])
    batch = review(client, headers, batch, decision="update", targetActivityId=before["id"], targetActivityRevision=before["revision"])
    result = commit(client, headers, batch)
    assert result["receipt"]["updated"] == 1
    after = client.get("/api/v1/friday-activities/" + before["id"], headers=headers).json()
    for field in ("ownerUserId", "description", "domain", "people"):
        assert after[field] == before[field]
    assert after["sessions"][0]["id"] == before["sessions"][0]["id"]
    assert after["sessions"][0]["location"] == "Updated room"
    assert after["sessions"][0]["attended"] == 3


def test_update_multi_session_target_is_rejected_during_review(import_client):
    client, headers = import_client
    target = create_target(client, headers, sessions=2)
    batch = preview(client, headers)
    batch = review(client, headers, batch, decision="update", targetActivityId=target["id"], targetActivityRevision=target["revision"])
    assert any("หนึ่งรอบ" in error for error in batch["rows"][0]["errors"])
    response = client.post(f"{PREFIX}/{batch['id']}/commit", headers=headers,
        json={"expectedBatchRevision": batch["revision"], "idempotencyKey": "preserve-two-sessions"})
    assert response.status_code == 400
    assert count("friday_sessions") == 2


@pytest.mark.parametrize("initial_status", ["contacting", "invited"])
def test_reviewed_historical_completion_follows_legal_states_for_existing_activity(import_client, initial_status):
    client, headers = import_client
    target = create_target(client, headers, status=initial_status)
    batch = preview(client, headers, [source_row(status="completed", attended="0")])
    batch = review(client, headers, batch, decision="update", targetActivityId=target["id"], targetActivityRevision=target["revision"])
    assert batch["rows"][0]["errors"] == []
    commit(client, headers, batch)
    after = client.get("/api/v1/friday-activities/" + target["id"], headers=headers).json()
    assert after["status"] == "completed"
    assert after["sessions"][0]["status"] == "held"
    assert after["sessions"][0]["attended"] == 0
    assert execute("SELECT count(*) FROM friday_audit_events WHERE operation='import.transition'", fetch=True) == [(3,)]


def test_reviewed_reopen_completed_target_keeps_held_history_and_replans_atomically(import_client):
    client, headers = import_client
    initial = review(client, headers, preview(client, headers, [source_row(status="completed", attended="7")]))
    first = commit(client, headers, initial)
    target_id = first["receipt"]["activityIds"][0]
    before = client.get("/api/v1/friday-activities/" + target_id, headers=headers).json()
    batch = preview(client, headers, [source_row(status="confirmed", sessionDate="2026-11-06", attended="")])
    batch = review(client, headers, batch, decision="update", targetActivityId=target_id, targetActivityRevision=before["revision"])
    assert batch["rows"][0]["errors"] == []
    result = commit(client, headers, batch)
    assert result["receipt"]["updated"] == 1
    after = client.get("/api/v1/friday-activities/" + target_id, headers=headers).json()
    assert after["status"] == "confirmed"
    assert after["sessions"][0]["status"] == "planned"
    assert after["sessions"][0]["sessionDate"] == "2026-11-06"
    assert after["sessions"][0]["id"] == before["sessions"][0]["id"]
    assert execute("""SELECT count(*) FROM friday_audit_events WHERE operation='import.transition'
        AND before_data->>'status'='completed' AND after_data->>'status'='confirmed' AND reason IS NOT NULL""", fetch=True) == [(1,)]
    assert execute("""SELECT count(*) FROM friday_audit_events WHERE after_data->>'status'='completed'
        AND (after_data->'sessions'->0->>'attended')::int=7""", fetch=True)[0][0] >= 1


def test_import_auth_rechecks_role_active_and_revoked_session(import_client):
    client, headers = import_client
    assert client.get(PREFIX + "/template.csv").status_code == 401
    execute("UPDATE users SET role=NULL WHERE id=1")
    assert client.post(PREFIX + "/preview", headers=headers, files={"file": ("a.csv", csv_bytes([source_row()]))}).status_code == 403
    execute("UPDATE users SET role='staff',is_active=false WHERE id=1")
    assert client.get(PREFIX + "/template.csv", headers=headers).status_code == 401
    execute("UPDATE users SET is_active=true WHERE id=1")
    client.post("/api/v1/auth/logout", headers=headers)
    assert client.get(PREFIX + "/template.csv", headers=headers).status_code == 401


def test_xlsx_native_date_time_and_strict_upload_limits(import_client, monkeypatch):
    from datetime import date, time
    client, headers = import_client
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(service.HEADERS))
    row = source_row()
    row.update(sessionDate=date(2026, 10, 16), startTime=time(13), endTime=time(16))
    sheet.append([row.get(header) for header in service.HEADERS])
    output = BytesIO()
    workbook.save(output)
    batch = preview(client, headers, content=output.getvalue(), name="synthetic.xlsx")
    assert batch["rows"][0]["errors"] == []
    assert batch["rows"][0]["normalized"]["startTime"] == "13:00"
    for content, name in ((b"bad", "bad.xlsx"), (b"\xff", "bad.csv"), (b"anything", "bad.xlsm")):
        response = client.post(PREFIX + "/preview", headers=headers, files={"file": (name, content)})
        assert response.status_code == 400
    monkeypatch.setattr(service, "MAX_FILE_BYTES", 10)
    assert client.post(PREFIX + "/preview", headers=headers, files={"file": ("large.csv", b"x" * 11)}).status_code == 400


def test_xlsx_bounded_expansion_entities_and_multiple_sheets_rejected():
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", b"x" * (service.MAX_EXPANDED_BYTES + 1))
    with pytest.raises(BadRequestError):
        service.parse_upload(output.getvalue(), "bomb.xlsx")
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("xl/workbook.xml", '<!DOCTYPE xml [<!ENTITY x "blocked">]><xml/>')
    with pytest.raises(BadRequestError):
        service.parse_upload(output.getvalue(), "entity.xlsx")
    workbook = Workbook()
    workbook.create_sheet("Extra")
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(BadRequestError):
        service.parse_upload(output.getvalue(), "multiple.xlsx")


@pytest.mark.parametrize("field,value", [("sessionDate", "2569-10-16"), ("academicYear", "2569"), ("endTime", "12:00"), ("capacity", "-1"), ("attended", "NaN"), ("format", "unknown source format"), ("status", "")])
def test_uncertain_values_never_silently_normalized_into_valid_facts(field, value):
    _, errors = service.normalize_row(source_row(**{field: value}))
    assert any(field in error for error in errors)
